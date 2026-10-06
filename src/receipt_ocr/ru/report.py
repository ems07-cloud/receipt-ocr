"""Папка с фото чеков → таблица Excel: каждый чек, итоги по месяцам, дубликаты и проблемы.

Главный источник — QR-код чека (точно и бесплатно). Если передан ``ReceiptProcessor`` из
основного пакета, нейросеть дополнительно:
  * распознаёт чеки, где QR не нашёлся (порван, засвечен), — с пометкой «по нейросети, проверить»;
  * сверяет свою сумму с суммой из QR — расхождение помечается, значит, одно из чтений ошибочно.
"""

import io
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Optional

from receipt_ocr.ru.fiscal import FiscalReceipt
from receipt_ocr.ru.qr import read_fiscal_qr

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
LLM_SCHEMA = {
    "merchant_name": "string",
    "transaction_date": "string (YYYY-MM-DD)",
    "transaction_time": "string (HH:MM)",
    "total_amount": "number",
}
MONTHS = [
    "Январь",
    "Февраль",
    "Март",
    "Апрель",
    "Май",
    "Июнь",
    "Июль",
    "Август",
    "Сентябрь",
    "Октябрь",
    "Ноябрь",
    "Декабрь",
]


@dataclass
class Entry:
    file: str
    receipt: Optional[FiscalReceipt] = None
    status: str = "ok"  # ok | duplicate | no_qr | llm_only | mismatch | error
    notes: list = field(default_factory=list)
    llm_total: Optional[Decimal] = None
    llm_when: Optional[datetime] = None
    merchant: str = ""

    @property
    def counted(self) -> bool:
        """Попадает ли чек в итоги: дубликаты и нераспознанные — нет."""
        return self.status in ("ok", "mismatch", "llm_only")

    @property
    def when(self) -> Optional[datetime]:
        return self.receipt.when if self.receipt else self.llm_when

    @property
    def amount(self) -> Optional[Decimal]:
        if self.receipt:
            return self.receipt.signed_total
        return self.llm_total


def _decimal(value) -> Optional[Decimal]:
    try:
        return (
            Decimal(str(value)).quantize(Decimal("0.01"))
            if value not in (None, "")
            else None
        )
    except InvalidOperation:
        return None


def _llm_when(data: dict) -> Optional[datetime]:
    date, time = (
        str(data.get("transaction_date") or ""),
        str(data.get("transaction_time") or "00:00"),
    )
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(f"{date} {time}".strip(), fmt)
        except ValueError:
            continue
    return None


def collect_images(paths: Iterable) -> list[Path]:
    out = []
    for p in map(Path, paths):
        if p.is_dir():
            out += sorted(x for x in p.rglob("*") if x.suffix.lower() in IMAGE_EXT)
        elif p.suffix.lower() in IMAGE_EXT:
            out.append(p)
    return out


def process_receipts(
    paths: Iterable, llm=None, model: Optional[str] = None
) -> list[Entry]:
    """Читает чеки. llm — необязательный ReceiptProcessor для нечитаемых QR и сверки сумм."""
    entries, seen = [], {}
    for path in collect_images(paths):
        e = Entry(file=path.name)
        try:
            e.receipt = read_fiscal_qr(str(path))
        except Exception as exc:  # битый файл не должен останавливать всю пачку
            e.status, e.notes = "error", [f"не открылся: {exc}"]
            entries.append(e)
            continue

        if llm is not None:
            try:
                data = llm.process_receipt(str(path), LLM_SCHEMA, model)
            except Exception as exc:
                data = {"error": str(exc)}
            if "error" not in data:
                e.llm_total, e.llm_when = (
                    _decimal(data.get("total_amount")),
                    _llm_when(data),
                )
                e.merchant = str(data.get("merchant_name") or "")

        if e.receipt is None:
            if e.llm_total is not None:
                e.status = "llm_only"
                e.notes.append("QR не прочитан — сумма по нейросети, проверьте")
            else:
                e.status = "no_qr"
                e.notes.append(
                    "QR-код не найден"
                    + ("" if llm else " — можно распознать нейросетью: --llm")
                )
        else:
            if e.receipt.key in seen:
                e.status = "duplicate"
                e.notes.append(f"тот же чек, что {seen[e.receipt.key]}")
            else:
                seen[e.receipt.key] = e.file
                if e.llm_total is not None and e.llm_total != e.receipt.total:
                    e.status = "mismatch"
                    e.notes.append(
                        f"нейросеть прочитала {e.llm_total}, а в QR {e.receipt.total}"
                    )
            if e.receipt.is_refund:
                e.notes.append(e.receipt.operation_name)
        entries.append(e)
    return entries


def monthly(entries: list[Entry]) -> "OrderedDict[tuple, dict]":
    """Итоги по месяцам: покупки, возвраты, итого и число чеков."""
    months: dict = {}
    for e in entries:
        if not e.counted or e.when is None or e.amount is None:
            continue
        m = months.setdefault(
            (e.when.year, e.when.month),
            {"buy": Decimal(0), "refund": Decimal(0), "n": 0},
        )
        m["n"] += 1
        if e.amount < 0:
            m["refund"] += -e.amount
        else:
            m["buy"] += e.amount
    return OrderedDict(sorted(months.items()))


STATUS_RU = {
    "ok": "",
    "duplicate": "дубликат — не учтён",
    "no_qr": "не распознан — не учтён",
    "llm_only": "по нейросети",
    "mismatch": "суммы расходятся",
    "error": "ошибка файла",
}


def build_excel(entries: list[Entry]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill, warn_fill = (
        PatternFill("solid", fgColor="1F4E78"),
        PatternFill("solid", fgColor="FFF2CC"),
    )
    bad_fill = PatternFill("solid", fgColor="F8D7DA")
    wb = Workbook()

    def header(ws, titles, widths):
        ws.append(titles)
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Чеки"
    header(
        ws,
        [
            "Дата и время",
            "Сумма, ₽",
            "Операция",
            "Продавец",
            "ФН",
            "ФД",
            "ФП",
            "Файл",
            "Статус",
            "Примечание",
        ],
        [17, 12, 16, 22, 19, 9, 13, 20, 22, 44],
    )
    for e in sorted(entries, key=lambda x: (x.when or datetime.max, x.file)):
        r = e.receipt
        ws.append(
            [
                e.when,
                float(e.amount) if e.amount is not None else None,
                r.operation_name if r else "",
                e.merchant,
                r.fn if r else "",
                r.fd if r else "",
                r.fp if r else "",
                e.file,
                STATUS_RU[e.status],
                "; ".join(e.notes),
            ]
        )
        row = ws.max_row
        ws.cell(row=row, column=1).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(row=row, column=2).number_format = "#,##0.00"
        fill = (
            bad_fill
            if e.status in ("duplicate", "no_qr", "error", "mismatch")
            else (warn_fill if e.status == "llm_only" else None)
        )
        if fill:
            for col in range(1, 11):
                ws.cell(row=row, column=col).fill = fill

    s = wb.create_sheet("По месяцам")
    header(
        s,
        ["Месяц", "Покупки, ₽", "Возвраты, ₽", "Итого, ₽", "Чеков"],
        [16, 14, 14, 14, 9],
    )
    total = {"buy": Decimal(0), "refund": Decimal(0), "n": 0}
    for (year, month), m in monthly(entries).items():
        s.append(
            [
                f"{MONTHS[month - 1]} {year}",
                float(m["buy"]),
                float(m["refund"]),
                float(m["buy"] - m["refund"]),
                m["n"],
            ]
        )
        for k in total:
            total[k] += m[k]
    s.append(
        [
            "Всего",
            float(total["buy"]),
            float(total["refund"]),
            float(total["buy"] - total["refund"]),
            total["n"],
        ]
    )
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    for row in s.iter_rows(min_row=2, min_col=2, max_col=4):
        for c in row:
            c.number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
